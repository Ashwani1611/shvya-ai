"""Workbook contracts for the tenant-authorised Message Activity read model."""

from contextlib import contextmanager
from datetime import datetime, timedelta, timezone as dt_timezone
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from uuid import UUID
from xml.etree import ElementTree
from zipfile import ZipFile, is_zipfile
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet import _writer
import pytest

from apps.superadmin import message_activity_export as exporter


UTC = dt_timezone.utc
RECORDED_AT = datetime(2026, 3, 8, 6, 45, tzinfo=UTC)


def activity_row(**changes):
    row = {
        "kind": "api_message",
        "kind_label": "WhatsApp message",
        "id": "activity-1",
        "event_time": RECORDED_AT,
        "origin": "AI reply",
        "status_label": "Delivered",
        "direction": "outbound",
        "channel": "WhatsApp API",
        "connection_label": "Sales number · +91 81691 95976",
        "lead_name": "Workbook Lead",
        "lead_phone": "+91 81691 95976",
        "content": "Here is the requested information.\nHow can we help?",
        "history": [],
        "related": [],
    }
    row.update(changes)
    return row


def export_rows(rows, *, tz="Asia/Kolkata", values=None, max_rows=50_000, target=None):
    target = BytesIO() if target is None else target
    count = exporter.write_activity_workbook(
        target,
        organization=SimpleNamespace(
            name="Workbook Organisation",
            pk=UUID("cc972246-21e1-43c1-8949-adc9bb7d5e80"),
        ),
        filters=SimpleNamespace(
            tz=ZoneInfo(tz),
            values=values or {"from": "2026-03-08", "to": "2026-03-08", "kind": ""},
        ),
        rows=rows,
        max_rows=max_rows,
    )
    assert not target.closed, "The export must leave its caller-owned target open."
    return target, count


@contextmanager
def read_workbook(target):
    target.seek(0)
    book = load_workbook(target, data_only=False)
    try:
        yield book
    finally:
        book.close()


def records(sheet):
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row, strict=True)) for row in sheet.iter_rows(min_row=2)]


def details(book):
    return {row[0].value: row[1].value for row in book["Export Details"].iter_rows(min_row=2)}


def test_export_is_real_xlsx_with_complete_activity_history_and_relationships():
    row = activity_row(
        campaign_name="Autumn campaign",
        history=[{"time": RECORDED_AT, "label": "Recorded", "detail": "Provider accepted message"}],
        related=[{"kind": "ai_reply", "id": "reply-1", "label": "Generated this message"}],
    )
    target, count = export_rows([row, activity_row(id="activity-2", content="Second record")])

    assert count == 2
    assert is_zipfile(target)
    target.seek(0)
    with ZipFile(target) as archive:
        assert "xl/workbook.xml" in archive.namelist()
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml" in (
            archive.read("[Content_Types].xml").decode()
        )

    with read_workbook(target) as book:
        assert book.sheetnames == ["Message Activity", "Recorded History", "Linked Records", "Export Details"]
        exported = records(book["Message Activity"])
        assert [row["Activity ID"].value for row in exported] == ["activity-1", "activity-2"]
        assert exported[0]["Message / content"].value == row["content"]
        assert exported[0]["Campaign"].value == "Autumn campaign"
        assert records(book["Recorded History"])[0]["Details"].value == "Provider accepted message"
        assert records(book["Linked Records"])[0]["Related ID"].value == "reply-1"
        assert book["Message Activity"].freeze_panes == "A2"
        assert book["Message Activity"].auto_filter.ref.endswith("3")
        assert details(book)["Organisation"] == "Workbook Organisation"
        assert details(book)["Exported activity records"] == 2
        assert details(book)["Timezone for date columns"] == "Asia/Kolkata"
        assert details(book)["Filter: from"] == "2026-03-08"


def test_phone_numbers_and_long_identifiers_are_exact_excel_text():
    long_id = 987654321098765432109876543210
    row = activity_row(
        id=long_id,
        lead_phone="+91 (81691) 95976",
        from_number=919988776655,
        to_number="0012345678901234567890",
        external_id=long_id,
        lead_id=long_id,
        connection_id=long_id,
        tokens=120,
        attempt=3,
        history=[{"time": RECORDED_AT, "label": "Recorded"}],
        related=[{"kind": "queue", "id": long_id, "label": "Queued by"}],
    )
    target, _ = export_rows([row])

    with read_workbook(target) as book:
        exported = records(book["Message Activity"])[0]
        expected = {
            "Lead number": "918169195976",
            "From number / ID": "919988776655",
            "To number / ID": "0012345678901234567890",
            "Provider message ID": str(long_id),
            "Activity ID": str(long_id),
            "Lead ID": str(long_id),
            "Connection ID": str(long_id),
        }
        for label, text in expected.items():
            assert exported[label].value == text, label
            assert exported[label].data_type == "s", label
            assert exported[label].number_format == "@", label
        for label, value in (("Attempt", 3), ("Recorded AI tokens", 120)):
            assert exported[label].value == value
            assert exported[label].data_type == "n"
        assert records(book["Recorded History"])[0]["Activity ID"].value == str(long_id)
        linked = records(book["Linked Records"])[0]
        assert linked["Activity ID"].value == str(long_id)
        assert linked["Related ID"].value == str(long_id)


@pytest.mark.parametrize("text", ["=2+2", "+SUM(1,2)", "-2+3", "@SUM(A1:A2)", '\t=HYPERLINK("https://example.invalid/","open")'])
def test_formula_like_untrusted_text_stays_literal_in_all_sheets(text):
    target, _ = export_rows(
        [activity_row(
            id=text,
            lead_name=text,
            content=text,
            error=text,
            sequence_name=text,
            history=[{"label": text, "detail": text}],
            related=[{"kind": text, "id": text, "label": text}],
        )],
        values={"q": text},
    )

    with read_workbook(target) as book:
        row = records(book["Message Activity"])[0]
        for label in ("Activity ID", "Lead", "Message / content", "Error / reason", "Sequence / Cadence"):
            assert row[label].value == text
            assert row[label].data_type == "s"
        assert records(book["Recorded History"])[0]["Details"].value == text
        assert records(book["Linked Records"])[0]["Relationship"].value == text
        assert details(book)["Filter: q"] == text
        assert all(cell.data_type != "f" for sheet in book for line in sheet for cell in line)

    target.seek(0)
    with ZipFile(target) as archive:
        for name in archive.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                xml = ElementTree.fromstring(archive.read(name))
                assert not xml.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f")


@pytest.mark.parametrize(
    ("tz", "expected_hours", "expected_minute"),
    [("Asia/Kolkata", [12, 13, 14, 15, 16, 17], 15), ("America/New_York", [1, 3, 4, 5, 6, 7], 45)],
)
def test_all_date_columns_and_history_use_organisation_timezone_including_dst(tz, expected_hours, expected_minute, monkeypatch):
    monkeypatch.setattr(exporter.timezone, "now", lambda: RECORDED_AT)
    keys = ["event_time", "scheduled_at", "started_at", "sent_at", "completed_at", "updated_at"]
    labels = ["Recorded at", "Scheduled for", "Processing started", "Provider sent at", "Processing completed", "Last record update"]
    dates = {key: RECORDED_AT + timedelta(hours=index) for index, key in enumerate(keys)}
    target, _ = export_rows(
        [activity_row(**dates, history=[{"time": dates["sent_at"], "label": "Provider accepted"}])],
        tz=tz,
    )

    with read_workbook(target) as book:
        row = records(book["Message Activity"])[0]
        for label, hour in zip(labels, expected_hours, strict=True):
            assert row[label].value == datetime(2026, 3, 8, hour, expected_minute), label
            assert row[label].data_type == "d", label
        history = records(book["Recorded History"])[0]
        assert history["Recorded at"].value == row["Provider sent at"].value
        assert details(book)["Timezone for date columns"] == tz
        assert details(book)["Generated at (UTC)"] == "2026-03-08T06:45:00+00:00"


def test_long_message_unicode_and_newlines_reconstruct_exactly_in_numbered_parts():
    content = '=message\n' + "🙂 café हिन्दी 漢字\r\n" * 2300
    error = "Delivery diagnostic\n" * 1800
    history_detail = "Retained history\n" * 2100
    target, _ = export_rows([activity_row(
        content=content,
        error=error,
        history=[{"time": RECORDED_AT, "label": "Recorded", "detail": history_detail}],
    )])

    with read_workbook(target) as book:
        assert "Long Content" in book.sheetnames
        exported = records(book["Message Activity"])[0]
        assert "Full text preserved" in exported["Message / content"].value
        assert len(exported["Message / content"].value) <= 16_000
        by_field = {}
        for row in records(book["Long Content"]):
            assert row["Activity"].value == "api_message"
            assert row["Activity ID"].value == "activity-1"
            text_cell = row["Text (join parts in order)"]
            assert text_cell.data_type == "s"
            assert len(text_cell.value.encode("utf-16-le")) // 2 <= 32_767
            by_field.setdefault(row["Field"].value, []).append((row["Part"].value, text_cell.value))
        reconstructed = {}
        for field, parts in by_field.items():
            assert [number for number, _ in parts] == list(range(1, len(parts) + 1))
            reconstructed[field] = "".join(text for _, text in parts)
        assert reconstructed["content"] == content
        assert reconstructed["error"] == error
        assert history_detail in reconstructed.values()


def test_invalid_xml_characters_are_removed_without_losing_valid_unicode_or_whitespace():
    invalid = "before\x00\x01\x08\x0b\x0c\x0e\x1f\ud800\udfff\ufffe\uffffafter\t\n🙂 café हिन्दी"
    expected = "beforeafter\t\n🙂 café हिन्दी"
    target, _ = export_rows([activity_row(content=invalid, error=invalid, lead_name=invalid)])

    with read_workbook(target) as book:
        row = records(book["Message Activity"])[0]
        for label in ("Message / content", "Error / reason", "Lead"):
            assert row[label].value == expected


def test_empty_export_is_readable_and_does_not_invent_activity_or_history():
    target, count = export_rows([])
    assert count == 0
    with read_workbook(target) as book:
        assert "Long Content" not in book.sheetnames
        for sheet in ("Message Activity", "Recorded History", "Linked Records"):
            assert records(book[sheet]) == []
            assert book[sheet].max_row == 1
        assert details(book)["Exported activity records"] == 0
        assert details(book)["Filter: kind"] == "All / not specified"


def test_export_limit_includes_exact_boundary_and_rejects_instead_of_returning_a_cut_file():
    target, count = export_rows([activity_row(id="1"), activity_row(id="2")], max_rows=2)
    assert count == 2
    with read_workbook(target) as book:
        assert len(records(book["Message Activity"])) == 2

    consumed = []

    def rows():
        for number in range(10):
            consumed.append(number)
            yield activity_row(id=str(number))

    failed_target = BytesIO()
    with pytest.raises(exporter.ExportTooLarge, match="More than 2 activities match"):
        export_rows(rows(), max_rows=2, target=failed_target)
    assert consumed == [0, 1, 2], "Read one extra row to prove overflow, then stop."
    assert not failed_target.closed
    assert failed_target.getvalue() == b"", "Never leave a truncated XLSX available as an export."


@pytest.mark.parametrize("sheet", ["Recorded History", "Linked Records", "Long Content"])
def test_auxiliary_sheet_row_limit_rejects_a_partial_workbook(monkeypatch, sheet):
    monkeypatch.setattr(exporter, "EXCEL_SHEET_ROW_LIMIT", 40)

    def row_with_parts(count):
        if sheet == "Recorded History":
            return activity_row(history=[{"label": "Recorded", "detail": str(i)} for i in range(count)])
        if sheet == "Linked Records":
            return activity_row(related=[{"kind": "ai", "id": str(i), "label": "Execution"} for i in range(count)])
        return activity_row(content="X" * (exporter.CELL_TEXT_LIMIT * count))

    target, count = export_rows([row_with_parts(39)])
    assert count == 1
    with read_workbook(target) as book:
        assert book[sheet].max_row == 40  # Header plus every preserved part.
    failed_target = BytesIO()
    with pytest.raises(exporter.ExportTooLarge, match=sheet):
        export_rows([row_with_parts(40)], target=failed_target)
    assert failed_target.getvalue() == b""
    assert not failed_target.closed


@pytest.fixture
def captured_xml_streams(monkeypatch):
    paths = []
    books = []
    original_temporary_file = _writer.create_temporary_file

    def temporary_file(*args, **kwargs):
        path = original_temporary_file(*args, **kwargs)
        paths.append(Path(path))
        return path

    def workbook(*args, **kwargs):
        book = Workbook(*args, **kwargs)
        books.append(book)
        return book

    monkeypatch.setattr(_writer, "create_temporary_file", temporary_file)
    monkeypatch.setattr(exporter, "Workbook", workbook)
    yield books, paths
    # Keep regression failures from leaving their own files or generators open.
    for book in books:
        for sheet in book.worksheets:
            if not sheet.closed:
                sheet.close()
            writer = getattr(sheet, "_writer", None)
            if writer and Path(writer.out).exists():
                writer.cleanup()
        book.close()


def assert_xml_streams_closed(captured):
    books, paths = captured
    assert books and paths
    assert all(book.write_only for book in books)
    assert all(sheet.closed for book in books for sheet in book.worksheets)
    assert not [str(path) for path in paths if path.exists()]


def test_write_only_xml_streams_are_cleaned_after_success(captured_xml_streams):
    export_rows([activity_row(content="Long content\n" * 3000)])
    assert_xml_streams_closed(captured_xml_streams)


@pytest.mark.parametrize("failure", ["overflow", "iteration", "save"])
def test_write_only_xml_streams_are_cleaned_on_export_failure(captured_xml_streams, monkeypatch, failure):
    def interrupted_rows():
        yield activity_row(content="Long content\n" * 3000)
        raise RuntimeError("Activity query was interrupted")

    def failed_save(self, target):
        raise OSError("Target write failed")

    if failure == "overflow":
        rows = [activity_row(id="1"), activity_row(id="2")]
        error, message, limit = exporter.ExportTooLarge, "More than 1 activities match", 1
    elif failure == "iteration":
        rows = interrupted_rows()
        error, message, limit = RuntimeError, "Activity query was interrupted", 5
    else:
        rows = [activity_row()]
        monkeypatch.setattr(Workbook, "save", failed_save)
        error, message, limit = OSError, "Target write failed", 5

    target = BytesIO()
    with pytest.raises(error, match=message):
        export_rows(rows, max_rows=limit, target=target)
    assert not target.closed
    assert_xml_streams_closed(captured_xml_streams)


def test_existing_xml_stream_is_cleaned_when_next_sheet_cannot_be_created(captured_xml_streams, monkeypatch):
    original_sheet = exporter._sheet
    calls = 0

    def failed_second_sheet(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise OSError("Temporary XML directory is full")
        return original_sheet(*args, **kwargs)

    monkeypatch.setattr(exporter, "_sheet", failed_second_sheet)
    with pytest.raises(OSError, match="Temporary XML directory is full"):
        export_rows([activity_row()])
    assert_xml_streams_closed(captured_xml_streams)
