"""Bounded, streaming XLSX export of the authorised Message Activity read model."""

from datetime import datetime
import re

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MAX_EXPORT_ROWS = 50_000
EXCEL_SHEET_ROW_LIMIT = 1_048_576
# Leave room for supplementary Unicode characters in Excel's UTF-16 cell limit.
CELL_TEXT_LIMIT = 16_000
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_PHONE = re.compile(r"[+\d\s().-]{7,}")

COLUMNS = (
    ("event_time", "Recorded at", 22),
    ("kind_label", "Activity", 22),
    ("origin", "Origin", 22),
    ("status_label", "Recorded status", 20),
    ("direction", "Direction", 14),
    ("channel", "Channel", 22),
    ("connection_label", "Linked number / account", 36),
    ("lead_name", "Lead", 26),
    ("lead_phone", "Lead number", 22),
    ("lead_email", "Email", 30),
    ("pipeline_name", "Current pipeline", 25),
    ("stage_name", "Current stage", 22),
    ("owner_name", "Pipeline owner", 24),
    ("sequence_name", "Sequence / Cadence", 28),
    ("campaign_name", "Campaign", 28),
    ("message_type", "Message type", 18),
    ("content", "Message / content", 65),
    ("error", "Error / reason", 48),
    ("from_number", "From number / ID", 25),
    ("to_number", "To number / ID", 25),
    ("scheduled_at", "Scheduled for", 22),
    ("started_at", "Processing started", 22),
    ("sent_at", "Provider sent at", 22),
    ("completed_at", "Processing completed", 22),
    ("updated_at", "Last record update", 22),
    ("attempt", "Attempt", 12),
    ("external_id", "Provider message ID", 40),
    ("model", "AI model", 25),
    ("tokens", "Recorded AI tokens", 20),
    ("id", "Activity ID", 42),
    ("lead_id", "Lead ID", 40),
    ("connection_id", "Connection ID", 44),
    ("content_coverage", "Content coverage", 40),
)


class ExportTooLarge(ValueError):
    """A full export needs narrower filters; never return a silently cut file."""


def phone_text(value):
    value = str(value or "")
    if _PHONE.fullmatch(value):
        return re.sub(r"\D", "", value)
    return value


def _text(value):
    return _INVALID_XML.sub("", str(value if value is not None else ""))


def _cell(sheet, value, *, tz, header=False):
    cell = WriteOnlyCell(sheet)
    if isinstance(value, datetime):
        cell.value = timezone.localtime(value, timezone=tz).replace(tzinfo=None) if timezone.is_aware(value) else value
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.value = value
    else:
        cell.value = _text(value)[:CELL_TEXT_LIMIT]
        # This is an actual XLSX string, including leading =,+,-,@. There are no
        # executable formulas and no scientific-notation or 15-digit ID loss.
        cell.data_type = "s"
        cell.number_format = "@"
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if header:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4338CA")
    return cell


def _sheet(workbook, name, columns, *, tz):
    sheet = workbook.create_sheet(name)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for index, (_, _, width) in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.append([_cell(sheet, label, tz=tz, header=True) for _, label, _ in columns])
    return sheet


def write_activity_workbook(target, *, organization, filters, rows, max_rows=MAX_EXPORT_ROWS):
    """Write every matching record in bounded batches; target stays caller-owned.

    Dates are real Excel date cells in the named organisation timezone. Phone
    numbers and identifiers are strings. Long text is preserved in ordered parts
    instead of relying on Excel's implicit 32,767-character truncation.
    """
    workbook = Workbook(write_only=True)
    try:
        workbook.properties.creator = "SHVYA AI"
        workbook.properties.title = "Organisation Message Activity"
        workbook.properties.subject = "Authorised organisation activity export"
        tz = filters.tz
        activity = _sheet(workbook, "Message Activity", COLUMNS, tz=tz)
        history = _sheet(workbook, "Recorded History", (
            ("kind", "Activity", 22), ("id", "Activity ID", 44),
            ("time", "Recorded at", 22), ("label", "Event", 28), ("detail", "Details", 70),
        ), tz=tz)
        links = _sheet(workbook, "Linked Records", (
            ("kind", "Activity", 22), ("id", "Activity ID", 44),
            ("related_kind", "Related activity", 22), ("related_id", "Related ID", 44),
            ("label", "Relationship", 60),
        ), tz=tz)
        details = _sheet(workbook, "Export Details", (("label", "Setting", 30), ("value", "Value", 110)), tz=tz)
        long_content = None
        history_count = links_count = long_count = count = 0
        sheet_rows = {}

        def append(sheet, cells):
            # Activity records are capped separately. One record can have many
            # history/link/content rows, which must also fit a valid workbook.
            row_count = sheet_rows.get(sheet.title, 1)  # Includes the header.
            if row_count >= EXCEL_SHEET_ROW_LIMIT:
                raise ExportTooLarge(
                    f"The {sheet.title} sheet exceeds Excel's row limit. "
                    "Narrow the dates or filters and export again."
                )
            sheet.append(cells)
            sheet_rows[sheet.title] = row_count + 1

        def append_text(sheet, values, *, kind="", record_id="", fields=()):
            nonlocal long_content, long_count
            safe_values = []
            for index, value in enumerate(values):
                if isinstance(value, str):
                    value = _text(value)
                    if len(value) > CELL_TEXT_LIMIT:
                        if long_content is None:
                            long_content = _sheet(workbook, "Long Content", (
                                ("kind", "Activity", 22), ("id", "Activity ID", 44),
                                ("field", "Field", 30), ("part", "Part", 10),
                                ("text", "Text (join parts in order)", 100),
                            ), tz=tz)
                        field = fields[index] if index < len(fields) else f"{sheet.title} column {index + 1}"
                        for part, offset in enumerate(range(0, len(value), CELL_TEXT_LIMIT), 1):
                            append(long_content, [_cell(long_content, item, tz=tz) for item in (
                                kind, record_id, field, part, value[offset:offset + CELL_TEXT_LIMIT],
                            )])
                            long_count += 1
                        value = value[:CELL_TEXT_LIMIT - 65] + "\n[Full text preserved in the Long Content sheet.]"
                safe_values.append(_cell(sheet, value, tz=tz))
            append(sheet, safe_values)

        for row in rows:
            count += 1
            if count > max_rows:
                raise ExportTooLarge(f"More than {max_rows:,} activities match. Narrow the dates or filters and export again.")
            data = dict(row)
            for key in ("lead_phone", "from_number", "to_number"):
                data[key] = phone_text(data.get(key))
            for key in ("id", "external_id", "lead_id", "connection_id"):
                data[key] = _text(data.get(key))
            data["content_coverage"] = "Stored preview; full text was not retained here" if row.get("content_is_preview") else "Stored message / activity content"
            append_text(activity, [data.get(key, "") for key, _, _ in COLUMNS],
                        kind=row["kind"], record_id=data["id"], fields=[key for key, _, _ in COLUMNS])
            for item in row.get("history", ()):
                append_text(history, [row["kind"], data["id"], item.get("time"), item.get("label", ""), item.get("detail", "")],
                            kind=row["kind"], record_id=data["id"])
                history_count += 1
            for item in row.get("related", ()):
                append_text(links, [row["kind"], data["id"], item.get("kind", ""), _text(item.get("id")), item.get("label", "")],
                            kind=row["kind"], record_id=data["id"])
                links_count += 1

        settings_rows = [
            ("Organisation", organization.name), ("Organisation ID", str(organization.pk)),
            ("Generated at (UTC)", timezone.now().isoformat()), ("Timezone for date columns", str(tz)),
            ("Exported activity records", count), ("Ordering", "Recorded time, activity type and ID, newest first"),
            ("Range meaning", "Recorded activity time in the organisation timezone; the end date is inclusive. Scheduled/sent/update times are separate columns."),
            ("Pipeline and stage", "Current CRM assignment. Any retained historical routing evidence appears in Recorded History."),
            ("Delivery meaning", "AI completion, queue completion, a Cadence attempt and a message are separate records. Sent means provider acceptance; delivered/read requires the saved message status."),
            ("History coverage", "Only retained records and timestamps are included. Last record update is not a delivery receipt timestamp. Deleted or uncaptured history cannot be reconstructed."),
            ("Scheduled activity", "Sequence entries are current active/paused schedule state, not an immutable record of every schedule change."),
            ("Phone and ID format", "Stored as Excel text. Phone country codes are preserved; a leading + and formatting punctuation are removed. No country code is guessed."),
            ("Long text", "Long Content, when present, preserves the full exported field in numbered parts. Join its parts in order. Invalid XML control characters are removed."),
            ("Sensitive diagnostics", "Message content is included for the authorised organisation. Raw provider payloads, credentials, prompts and internal AI reasoning are excluded."),
        ]
        settings_rows.extend((f"Filter: {key}", value or "All / not specified") for key, value in filters.values.items())
        for pair in settings_rows:
            append_text(details, pair)
        for sheet, row_count, column_count in (
            (activity, count, len(COLUMNS)), (history, history_count, 5), (links, links_count, 5),
        ):
            sheet.auto_filter.ref = f"A1:{get_column_letter(column_count)}{row_count + 1}"
        if long_content is not None:
            long_content.auto_filter.ref = f"A1:E{long_count + 1}"
        workbook.save(target)
        return count
    finally:
        # openpyxl creates temporary XML streams for write-only sheets. Close and
        # remove them even if the query/export limit fails before save completes.
        for sheet in workbook.worksheets:
            if not sheet.closed:
                sheet.close()
            writer = getattr(sheet, "_writer", None)
            if writer and getattr(writer, "out", None):
                try:
                    writer.cleanup()
                except FileNotFoundError:
                    pass
        workbook.close()
