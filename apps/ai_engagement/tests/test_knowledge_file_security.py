from __future__ import annotations

import io
import zipfile
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from docx import Document as DocxDocument
from openpyxl import Workbook
from pypdf import PdfWriter

from apps.ai_engagement.models import Document
from apps.ai_engagement.serializers.document import DocumentUploadSerializer
from apps.ai_engagement.services.knowledge_file_security import (
    KnowledgeFileSecurityError,
    validate_knowledge_file,
    validate_organization_knowledge_quota,
)
from apps.ai_engagement.services.knowledge_source import (
    KnowledgeSourceService,
    KnowledgeSourceServiceError,
)
from apps.organizations.models import Organization


class KnowledgeFileValidationTests(SimpleTestCase):
    def test_accepts_valid_utf8_text(self):
        uploaded = SimpleUploadedFile(
            "guide.txt",
            b"SHVYA knowledge content",
            content_type="application/octet-stream",
        )

        inspection = validate_knowledge_file(
            uploaded,
        )

        self.assertEqual(inspection.extension, ".txt")
        self.assertEqual(inspection.size, len(b"SHVYA knowledge content"))
        self.assertEqual(uploaded.tell(), 0)

    def test_serializer_rejects_actual_stream_over_size_limit(self):
        uploaded = SimpleUploadedFile(
            "large.txt",
            b"a" * 33,
            content_type="text/plain",
        )

        with patch(
            "apps.ai_engagement.services.knowledge_file_security."
            "MAX_UPLOAD_BYTES",
            32,
        ):
            serializer = DocumentUploadSerializer(
                data={"file": uploaded}
            )

            self.assertFalse(serializer.is_valid())

        self.assertIn("file", serializer.errors)
        self.assertIn(
            "too large",
            str(serializer.errors["file"][0]).lower(),
        )

    def test_rejects_binary_data_disguised_as_text(self):
        uploaded = SimpleUploadedFile(
            "notes.txt",
            b"safe-prefix\x00binary-data",
            content_type="text/plain",
        )

        with self.assertRaisesRegex(
            KnowledgeFileSecurityError,
            "binary null bytes",
        ):
            validate_knowledge_file(uploaded)

    def test_rejects_invalid_utf8_text(self):
        uploaded = SimpleUploadedFile(
            "notes.csv",
            b"name,value\nhello,\xff",
            content_type="text/csv",
        )

        with self.assertRaisesRegex(
            KnowledgeFileSecurityError,
            "valid UTF-8",
        ):
            validate_knowledge_file(uploaded)

    def test_rejects_plain_text_renamed_to_pdf(self):
        uploaded = SimpleUploadedFile(
            "fake.pdf",
            b"This is not a PDF",
            content_type="application/pdf",
        )

        with self.assertRaisesRegex(
            KnowledgeFileSecurityError,
            "does not match the .pdf extension",
        ):
            validate_knowledge_file(uploaded)

    def test_accepts_real_pdf_even_with_spoofed_content_type(self):
        buffer = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(buffer)

        uploaded = SimpleUploadedFile(
            "guide.pdf",
            buffer.getvalue(),
            content_type="text/plain",
        )

        inspection = validate_knowledge_file(uploaded)

        self.assertEqual(inspection.extension, ".pdf")

    def test_rejects_pdf_over_page_limit(self):
        buffer = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.add_blank_page(width=100, height=100)
        writer.write(buffer)

        uploaded = SimpleUploadedFile(
            "guide.pdf",
            buffer.getvalue(),
            content_type="application/pdf",
        )

        with patch(
            "apps.ai_engagement.services.knowledge_file_security."
            "MAX_PDF_PAGES",
            1,
        ):
            with self.assertRaisesRegex(
                KnowledgeFileSecurityError,
                "too many pages",
            ):
                validate_knowledge_file(uploaded)

    def test_accepts_real_docx(self):
        buffer = io.BytesIO()
        document = DocxDocument()
        document.add_paragraph("Safe knowledge")
        document.save(buffer)

        uploaded = SimpleUploadedFile(
            "guide.docx",
            buffer.getvalue(),
            content_type="application/zip",
        )

        inspection = validate_knowledge_file(uploaded)

        self.assertEqual(inspection.extension, ".docx")

    def test_accepts_real_xlsx(self):
        buffer = io.BytesIO()
        workbook = Workbook()
        workbook.active["A1"] = "Safe knowledge"
        workbook.save(buffer)
        workbook.close()

        uploaded = SimpleUploadedFile(
            "guide.xlsx",
            buffer.getvalue(),
            content_type="application/zip",
        )

        inspection = validate_knowledge_file(uploaded)

        self.assertEqual(inspection.extension, ".xlsx")

    def test_rejects_zip_renamed_to_docx_without_docx_structure(self):
        buffer = io.BytesIO()

        with zipfile.ZipFile(
            buffer,
            mode="w",
            compression=zipfile.ZIP_DEFLATED,
        ) as archive:
            archive.writestr("random.txt", "not a docx")

        uploaded = SimpleUploadedFile(
            "fake.docx",
            buffer.getvalue(),
            content_type="application/zip",
        )

        with self.assertRaisesRegex(
            KnowledgeFileSecurityError,
            "does not match a valid .docx",
        ):
            validate_knowledge_file(uploaded)

    def test_rejects_office_archive_expanding_past_limit(self):
        buffer = io.BytesIO()

        with zipfile.ZipFile(
            buffer,
            mode="w",
            compression=zipfile.ZIP_DEFLATED,
        ) as archive:
            archive.writestr("[Content_Types].xml", "types")
            archive.writestr("word/document.xml", "document")
            archive.writestr("word/large.xml", "a" * 2_048)

        uploaded = SimpleUploadedFile(
            "bomb.docx",
            buffer.getvalue(),
            content_type="application/zip",
        )

        with patch(
            "apps.ai_engagement.services.knowledge_file_security."
            "MAX_ARCHIVE_UNCOMPRESSED_BYTES",
            1_024,
        ):
            with self.assertRaisesRegex(
                KnowledgeFileSecurityError,
                "decompression limit",
            ):
                validate_knowledge_file(uploaded)

    def test_rejects_office_archive_with_unsafe_path(self):
        buffer = io.BytesIO()

        with zipfile.ZipFile(
            buffer,
            mode="w",
            compression=zipfile.ZIP_DEFLATED,
        ) as archive:
            archive.writestr("[Content_Types].xml", "types")
            archive.writestr("word/document.xml", "document")
            archive.writestr("../outside.xml", "unsafe")

        uploaded = SimpleUploadedFile(
            "unsafe.docx",
            buffer.getvalue(),
            content_type="application/zip",
        )

        with self.assertRaisesRegex(
            KnowledgeFileSecurityError,
            "unsafe archive path",
        ):
            validate_knowledge_file(uploaded)

    def test_rejects_suspicious_office_compression_ratio(self):
        buffer = io.BytesIO()

        with zipfile.ZipFile(
            buffer,
            mode="w",
            compression=zipfile.ZIP_DEFLATED,
        ) as archive:
            archive.writestr("[Content_Types].xml", "types")
            archive.writestr("word/document.xml", "document")
            archive.writestr("word/repeated.xml", "a" * (1024 * 1024))

        uploaded = SimpleUploadedFile(
            "compressed.docx",
            buffer.getvalue(),
            content_type="application/zip",
        )

        with patch(
            "apps.ai_engagement.services.knowledge_file_security."
            "MAX_ARCHIVE_COMPRESSION_RATIO",
            2,
        ), patch(
            "apps.ai_engagement.services.knowledge_file_security."
            "MAX_ARCHIVE_UNCOMPRESSED_BYTES",
            4 * 1024 * 1024,
        ):
            with self.assertRaisesRegex(
                KnowledgeFileSecurityError,
                "suspiciously compressed",
            ):
                validate_knowledge_file(uploaded)

    def test_rejects_malformed_office_zip(self):
        uploaded = SimpleUploadedFile(
            "broken.docx",
            b"PK\x03\x04not-a-real-zip",
            content_type="application/zip",
        )

        with self.assertRaisesRegex(
            KnowledgeFileSecurityError,
            "not a valid Office archive",
        ):
            validate_knowledge_file(uploaded)


class KnowledgeFileServiceSecurityTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.organization = Organization.objects.create(
            name="Knowledge File Security Organization",
        )

    def test_file_source_rejects_content_extension_mismatch(self):
        uploaded = SimpleUploadedFile(
            "fake.pdf",
            b"not a pdf",
            content_type="application/pdf",
        )

        service = KnowledgeSourceService()

        with self.assertRaisesRegex(
            KnowledgeSourceServiceError,
            "does not match the .pdf extension",
        ):
            service.create_file_source(
                organization=self.organization,
                uploaded_file=uploaded,
            )

        self.assertFalse(
            Document.objects.filter(
                organization=self.organization,
            ).exists()
        )

    def test_organization_storage_quota_counts_existing_files(self):
        existing = SimpleUploadedFile(
            "existing.txt",
            b"12345",
            content_type="text/plain",
        )

        Document.objects.create(
            organization=self.organization,
            name="Existing",
            source_key="existing.txt",
            version=1,
            file=existing,
            processing_status=Document.ProcessingStatus.COMPLETED,
            is_active=True,
        )

        with patch(
            "apps.ai_engagement.services.knowledge_file_security."
            "MAX_ORGANIZATION_STORAGE_BYTES",
            8,
        ):
            with self.assertRaisesRegex(
                KnowledgeFileSecurityError,
                "quota exceeded",
            ):
                validate_organization_knowledge_quota(
                    organization=self.organization,
                    incoming_size=4,
                )

    def test_service_enforces_organization_storage_quota(self):
        existing = SimpleUploadedFile(
            "existing-service.txt",
            b"12345",
            content_type="text/plain",
        )

        Document.objects.create(
            organization=self.organization,
            name="Existing service file",
            source_key="existing-service.txt",
            version=1,
            file=existing,
            processing_status=Document.ProcessingStatus.COMPLETED,
            is_active=True,
        )

        incoming = SimpleUploadedFile(
            "incoming.txt",
            b"6789",
            content_type="text/plain",
        )

        service = KnowledgeSourceService()

        with patch(
            "apps.ai_engagement.services.knowledge_file_security."
            "MAX_ORGANIZATION_STORAGE_BYTES",
            8,
        ):
            with self.assertRaisesRegex(
                KnowledgeSourceServiceError,
                "quota exceeded",
            ):
                service.create_file_source(
                    organization=self.organization,
                    uploaded_file=incoming,
                )
